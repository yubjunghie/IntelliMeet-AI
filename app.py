import streamlit as st
import streamlit.components.v1 as components
from datetime import datetime
import os
from dotenv import load_dotenv
import google.generativeai as genai
from streamlit_mic_recorder import mic_recorder
import tempfile
import time
import io
import requests
from bs4 import BeautifulSoup
from docx import Document
import hmac
import markdown
import urllib.parse
import re

# 세션 상태 초기화
if 'first_pass_result' not in st.session_state:
    st.session_state.first_pass_result = None
if 'post_process_result' not in st.session_state:
    st.session_state.post_process_result = None

# 1. 환경 변수(.env) 또는 클라우드(Secrets)에서 API 키 불러오기
if "GEMINI_API_KEY" in st.secrets:
    GOOGLE_API_KEY = st.secrets["GEMINI_API_KEY"]
else:
    load_dotenv(override=True)
    GOOGLE_API_KEY = os.getenv("GEMINI_API_KEY")

if GOOGLE_API_KEY and GOOGLE_API_KEY != "여기에_발급받은_API_키를_입력하세요":
    genai.configure(api_key=GOOGLE_API_KEY)
    # [수정] 오디오 분석 및 화자 분리(Diarization) 성능이 월등히 뛰어난 최신 gemini-1.5-pro 모델로 변경합니다.
    model = genai.GenerativeModel('gemini-1.5-pro')
else:
    model = None

# 페이지 기본 설정
st.set_page_config(page_title="AI 회의록 작성 프로그램", page_icon="📝", layout="wide")

def check_password():
    def password_entered():
        if hmac.compare_digest(st.session_state["password"], st.secrets["password"]):
            st.session_state["password_correct"] = True
            del st.session_state["password"]
        else:
            st.session_state["password_correct"] = False

    if st.session_state.get("password_correct", False):
        return True

    st.title("🔒 로그인")
    st.text_input("프로그램 접속 비밀번호를 입력하세요:", type="password", on_change=password_entered, key="password")
    if "password_correct" in st.session_state:
        st.error("😕 비밀번호가 틀렸습니다. 다시 시도해 주세요.")
    return False

if not check_password():
    st.stop()

st.markdown("""
<style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    
    /* 전체 레이아웃 및 폰트 개선 */
    @import url('https://fonts.googleapis.com/css2?family=Pretendard:wght@400;500;600;700;900&display=swap');
    html, body, [class*="css"] {
        font-family: 'Pretendard', sans-serif;
    }

    .main-title {
        font-size: 2.8rem;
        font-weight: 900;
        background: -webkit-linear-gradient(45deg, #1e3c72, #2a5298);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 0px;
        text-align: center;
        padding-top: 10px;
    }
    .sub-title {
        font-size: 1.1rem;
        color: #555555;
        text-align: center;
        margin-bottom: 5px;
        font-weight: 500;
    }

    /* 모바일 UI (스마트폰 최적화) */
    @media (max-width: 768px) {
        .main-title {
            font-size: 2.0rem;
            padding-top: 5px;
        }
        .sub-title {
            font-size: 0.95rem;
        }
        /* 컨테이너 패딩 축소 */
        .block-container {
            padding-top: 1.5rem !important;
            padding-left: 1rem !important;
            padding-right: 1rem !important;
        }
        /* 입력 폼 여백 최적화 */
        .stTextInput, .stSelectbox, .stDateInput, .stTimeInput {
            margin-bottom: 0.5rem;
        }
    }
</style>
""", unsafe_allow_html=True)

st.markdown("<h1 class='main-title'>✨ IntelliMeet AI</h1>", unsafe_allow_html=True)
st.markdown("<p class='sub-title'>스마트한 비즈니스를 위한 AI 회의록 작성 솔루션</p>", unsafe_allow_html=True)
st.markdown("<div style='text-align: center; margin-bottom: 25px; color: #888; font-size: 0.9rem;'><strong>개발:</strong> BY CDY &nbsp;|&nbsp; <strong>문의:</strong> chady@ktc.re.kr</div>", unsafe_allow_html=True)



if model is None:
    st.error("⚠️ `.env` 파일에 Google Gemini API 키가 설정되지 않았습니다.")
    st.stop()

# [수정] 모바일 앱(클로바노트)처럼 보이기 위해 2단 레이아웃을 없애고 마이크 기능을 중앙으로 가져왔습니다.

# 🎙️ 메인 기능 (녹음 및 업로드)
st.markdown("### 🎙️ 새 회의 기록하기")
st.markdown("음성을 녹음하거나 파일을 업로드하여 AI 회의록을 자동 생성하세요.")

with st.container(border=True):
    tab_record, tab_upload = st.tabs(["⏺️ 직접 녹음", "📁 파일 업로드"])
    
    with tab_record:
        st.markdown("<div style='text-align: center; padding: 20px; color: #555;'>아래 버튼을 눌러 회의 녹음을 시작하세요.</div>", unsafe_allow_html=True)
        # 마이크 버튼을 화면 정중앙에 예쁘게 배치하기 위해 빈 컬럼을 양옆에 둡니다.
        col_mic1, col_mic2, col_mic3 = st.columns([1, 2, 1])
        with col_mic2:
            audio = mic_recorder(
                start_prompt="🎙️ AI 마이크 켜기",
                stop_prompt="⏹️ AI 마이크 끄기 (녹음 완료)",
                just_once=True,
                key='recorder'
            )
            
    with tab_upload:
        st.info("기존에 녹음된 회의 음성 파일(mp3, wav, m4a 등)을 업로드하세요.")
        uploaded_audio = st.file_uploader("음성 파일 선택", type=['wav', 'mp3', 'm4a', 'ogg', 'flac'])

st.write("")

# ⚙️ 상세 설정 (기본적으로 접혀있어서 화면을 가리지 않음)
with st.expander("⚙️ 회의 상세 설정 및 참고 자료 (선택)"):
    st.markdown("입력하지 않으면 오늘 날짜를 기준으로 기본값이 자동 설정됩니다.")
    
    # 기본 행사 제목을 오늘 날짜로 자동 설정
    today_str = datetime.today().strftime('%Y년 %m월 %d일')
    event_title = st.text_input("행사 제목", f"{today_str} 회의")
    
    col1, col2 = st.columns(2)
    with col1:
        event_type = st.selectbox("행사 유형", ("회의", "세미나", "교육", "포럼", "과제평가", "심사원 과정 교육 (시험대비)"))
    with col2:
        event_date = st.date_input("날짜", datetime.today())
        event_time = st.time_input("시간", datetime.now().time())
        
    col3, col4 = st.columns(2)
    with col3:
        role_type = st.selectbox("담당자 역할", ("참석자", "강의자", "발표자", "평가자"))
    with col4:
        person_name = st.text_input("담당자 이름", "홍길동")
        
    event_lang = st.radio("회의 진행 언어", ("한국어", "영어 (영문 원본 및 한글 번역본 동시 제공)"), horizontal=True)
    
    st.markdown("---")
    st.markdown("**📚 행사 커리큘럼 (안건)**")
    curriculums = {
        "회의": "1. 개회 및 참석자 소개 (회의 목적, 이전 회의 Follow-up)\n2. 핵심 현안 논의 (상황 분석, 부서별 의견 및 데이터 공유)\n3. 대안 모색 및 해결 방안 도출 (브레인스토밍, 리스크 검토)\n4. 향후 계획 수립 및 역할 분담 (Action Plan, 담당자 및 기한 설정)\n5. 요약 및 폐회 (결정 사항 최종 확인)",
        "세미나": "1. 연사 소개 및 배경 설명 (발표자 이력, 세미나 개최 취지)\n2. 메인 주제 심층 발표 (최신 동향, 연구 결과, 사례 분석)\n3. 전문가 패널 토의 (다각적 관점의 이슈 분석 및 논쟁점)\n4. 청중 참여 Q&A (질의응답 및 추가 해설)\n5. Wrap-up 및 네트워킹 안내",
        "교육": "1. 오리엔테이션 및 학습 목표 (과정 개요, 성취 기대치)\n2. 핵심 이론 및 개념 강의 (주요 원리, 표준 프로세스 설명)\n3. 실무 적용 및 실습 (사례 연구, 그룹 액티비티, 시뮬레이션)\n4. 성과 측정 및 평가 (이해도 점검, 퀴즈, 피드백 제공)\n5. 교육 내용 총정리 및 향후 학습 가이드",
        "포럼": "1. 오프닝 및 기조 연설 (행사 비전, 인사말)\n2. 세션별 발제자 주제 발표 (주요 의제별 심층 발제)\n3. 지정 토론 및 패널 디스커션 (이슈별 찬반 및 대안 논의)\n4. 청중 자유 토론 및 의견 수렴 (오픈 마이크)\n5. 종합 요약, 선언문 채택 및 폐회",
        "과제평가": "1. 과제 추진 배경 및 목표 소개 (과제 개요)\n2. 성과 및 실적 발표 (주요 산출물, 목표 달성률, 예산 집행 내역)\n3. 평가 위원 질의응답 (기술적 한계, 문제 해결 과정 검증)\n4. 보완/개선 사항 피드백 (평가위원 종합 의견)\n5. 최종 평가 점수 산정 및 향후 조치사항 정리",
        "심사원 과정 교육 (시험대비)": "1. ISO/국제 표준 규격 핵심 요구사항 심층 해설 (조항별 팩트 체크)\n2. 부적합 사례 및 심사 기법 연구 (실무 적용 포인트)\n3. 모의 심사 롤플레이 및 강사 피드백 (실전 감각 배양)\n4. ★시험 대비 핵심 요약 (자주 출제되는 개념 정리)\n5. 예상 문제 풀이 및 오답 노트 (Q&A 포함)"
    }
    default_curriculum = curriculums.get(event_type, "")
    curriculum_text = st.text_area("AI 분석에 반영될 진행 순서입니다. 필요에 따라 수정하세요.", value=default_curriculum, height=150, key=f"curriculum_{event_type}")

    st.markdown("---")
    st.markdown("**📎 참고 자료 첨부 (선택 사항)**")
    context_file = st.file_uploader("참고 문서 업로드", type=['pdf', 'txt', 'md', 'csv', 'docx'])
    context_url = st.text_input("참고 웹사이트 주소 (URL)")

audio_bytes = None
if uploaded_audio is not None:
    audio_bytes = uploaded_audio.read()
elif audio is not None:
    audio_bytes = audio['bytes']

if audio_bytes:
    st.audio(audio_bytes, format="audio/webm")
    st.success("음성 데이터가 성공적으로 인식되었습니다! AI가 분석을 시작합니다... (잠시만 기다려 주세요 ⏳)")
    
    with tempfile.NamedTemporaryFile(delete=False, suffix=".webm") as tmp_file:
        tmp_file.write(audio_bytes)
        tmp_file_path = tmp_file.name

    try:
        with st.spinner('인공지능이 음성을 분석하고 회의록을 작성 중입니다...'):
            files_to_send = []
            
            if context_file:
                file_ext = context_file.name.split('.')[-1] if '.' in context_file.name else 'txt'
                with tempfile.NamedTemporaryFile(delete=False, suffix=f".{file_ext}") as tmp_ctx:
                    tmp_ctx.write(context_file.read())
                    tmp_ctx_path = tmp_ctx.name
                uploaded_ctx = genai.upload_file(path=tmp_ctx_path)
                files_to_send.append(uploaded_ctx)

            if context_url:
                try:
                    res = requests.get(context_url, timeout=5)
                    soup = BeautifulSoup(res.text, 'html.parser')
                    url_text = soup.get_text(separator=' ', strip=True)
                    prompt_url = f"\n\n**[웹사이트 참고 자료]** 다음은 첨부된 웹사이트({context_url})의 텍스트입니다:\n{url_text[:5000]}\n"
                    files_to_send.append(prompt_url)
                except Exception as e:
                    st.warning(f"웹사이트 주소를 불러오는데 실패했습니다: {e}")

            uploaded_file = genai.upload_file(path=tmp_file_path)
            files_to_send.append(uploaded_file)
            
            # [수정] 클로바노트처럼 똑똑한 요약을 위해 프롬프트(AI 지시사항)를 아주 상세하게 변경했습니다.
            # SECTION_DIVIDER를 추가하여 나중에 탭(Tab) 화면으로 쉽게 쪼갤 수 있게 만듭니다.
            prompt = f"""
당신은 최고 수준의 AI 회의록 분석 비서입니다. 첨부된 오디오 파일을 깊이 있게 분석하여, '클로바노트(Clova Note)' 앱처럼 스마트하고 체계적인 회의록을 작성해 주세요.
진행된 행사의 제목은 '{event_title}', 종류는 '{event_type}'이며, 일시는 {event_date} {event_time} 입니다.
{role_type}는 {person_name}님 입니다.

**[진행 순서 및 안건]**
{curriculum_text}
위의 안건을 참고하여, 전체 대화의 맥락을 정확하게 파악해 주세요.
"""
            if context_file:
                prompt += "\n**[중요] 오디오 파일과 함께 참고 자료가 첨부되었습니다. 이 참고 자료의 내용을 바탕으로 전문 용어나 고유 명사를 정확하게 인식하고, 대화의 문맥을 더 정확히 분석해 주세요.**\n"
                
            if "영어" in event_lang:
                prompt += "\n**[중요 언어 지침] 이 회의는 주로 영어로 진행되었습니다. 모든 분석 결과와 대화록은 먼저 '영어(English)' 원문으로 상세히 작성하고, 각 항목 바로 아래에 '한국어 번역(Korean Translation)'을 병기해 주세요.**\n"

            prompt += f"""
[작성 규칙 및 출력 형식]
아래 지정된 [구분선] `<!-- SECTION_DIVIDER -->` 을 사용하여 텍스트를 정확히 3개의 파트로 나누어 출력해 주세요. 이 구분선은 모바일 앱 화면에서 탭(Tab)으로 나누어 보여주기 위해 필수적입니다.
아래의 마크다운 구조를 엄격하게 지켜주세요.

# 📝 [{event_type}] {event_title} 요약 리포트

**일시:** {event_date} {event_time}
**{role_type}:** {person_name}

---

### 🌟 1. 한 줄 핵심 요약 (TL;DR)
(전체 회의의 가장 중요한 핵심이나 결론을 2~3줄로 매우 간결하게 요약)

### ⏱️ 2. 시간대별/주제별 주요 논의 흐름 (Timeline Summary)
(회의의 진행 흐름에 따라 안건이나 주제가 넘어가는 지점을 포착하여, '어떤 주제로 넘어가서 무슨 이야기가 오갔는지'를 흐름대로 요약)
- **도입부:** ...
- **주요 논의 1:** ...
- **주요 논의 2:** ...
- **마무리:** ...
"""
            if event_type == "심사원 과정 교육 (시험대비)":
                prompt += "\n**[특수 조건] 흐름 요약 부분에 강사가 강조한 부분과 시험 예상 문제/힌트, ISO 9001 주요 심사 포인트를 눈에 띄게 강조해 주세요.**\n"
            elif event_type == "과제평가":
                prompt += "\n**[특수 조건] 흐름 요약 부분에 평가 위원들의 날카로운 지적 사항과 피드백, 기술적 한계에 대한 평가를 중점적으로 기록해 주세요.**\n"

            prompt += f"""
### 👥 3. 참석자(화자)별 주요 의견 요약
(Speaker A, Speaker B 등 화자별로 나누어, 각 인물이 주로 어떤 주장을 했고 어떤 의견을 내었는지 요약)
- **참석자 A:** ...
- **참석자 B:** ...

### ✅ 4. 다음 할 일 및 결정 사항 (Action Items & Decisions)
(최종적으로 합의된 사항과 앞으로 누가 무엇을 언제까지 해야 하는지 표로 명확히 정리)
| 담당자 (화자) | 할 일 (Task) | 기한/비고 |
|---|---|---|
| ... | ... | ... |

<!-- SECTION_DIVIDER -->

### 🗣️ 5. 화자가 분리된 전체 대화록 (Full Transcript)
(전체 대화 내용을 누가 말했는지 명확히 분리하여 [참석자 A], [참석자 B] 와 같이 말머리를 달아 스크립트로 작성)
[참석자 A]: 안녕하세요, 회의 시작하겠습니다.
[참석자 B]: 네, 첫 번째 안건부터 보시죠...

<!-- SECTION_DIVIDER -->

### 🧠 6. 회의 마인드맵 (Mind Map)
(회의 내용의 핵심 키워드들을 계층 구조로 정리하여 아래와 같은 Mermaid js 문법으로 출력. 반드시 ````mermaid 로 시작하고 ```` 로 끝나야 함)
````mermaid
mindmap
  root((회의 주제))
    키워드1
      상세1
      상세2
    키워드2
      상세3
````
"""
            
            files_to_send.append(prompt)
            response = model.generate_content(files_to_send)
            
            st.session_state.first_pass_result = response.text
            st.session_state.post_process_result = None
            
            genai.delete_file(uploaded_file.name)
            if context_file:
                genai.delete_file(uploaded_ctx.name)
            
    except Exception as e:
        st.error(f"오류가 발생했습니다: {e}")
    finally:
        if os.path.exists(tmp_file_path):
            os.remove(tmp_file_path)
        if context_file and os.path.exists(tmp_ctx_path):
            os.remove(tmp_ctx_path)

if st.session_state.first_pass_result:
    st.divider()
    st.subheader("✨ AI 회의록 분석 결과")
    
    display_text = st.session_state.post_process_result if st.session_state.post_process_result else st.session_state.first_pass_result
    
    # [수정] 모바일 어플리케이션처럼 보이도록 화면을 3개의 탭(Tab)으로 예쁘게 나눕니다.
    # 앞서 프롬프트에서 지정한 <!-- SECTION_DIVIDER --> 를 기준으로 텍스트를 쪼갭니다.
    sections = display_text.split("<!-- SECTION_DIVIDER -->")
    
    summary_text = sections[0].strip() if len(sections) > 0 else display_text
    transcript_text = sections[1].strip() if len(sections) > 1 else "대화록을 분리하지 못했습니다. 요약 탭을 확인해 주세요."
    mindmap_text = sections[2].strip() if len(sections) > 2 else ""

    # 클로바노트 스타일의 3가지 탭 생성
    tab_summary, tab_transcript, tab_mindmap = st.tabs(["📑 요약 노트", "🗣️ 전체 대화록", "🧠 마인드맵"])
    
    with tab_summary:
        with st.container(border=True):
            st.markdown(summary_text)
            
    with tab_transcript:
        with st.container(border=True):
            st.markdown(transcript_text)
            
    with tab_mindmap:
        # 마인드맵 렌더링
        mermaid_match = re.search(r'```mermaid\n(.*?)\n```', display_text, re.DOTALL)
        if mermaid_match:
            with st.container(border=True):
                st.info("💡 모바일 기기에서는 이미지를 터치하고 드래그하여 움직일 수 있습니다.")
                mermaid_code = mermaid_match.group(1).strip()
                htmlcode = f"""
                <div class="mermaid">
                {mermaid_code}
                </div>
                <script type="module">
                  import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.esm.min.mjs';
                  mermaid.initialize({{ startOnLoad: true }});
                </script>
                """
                components.html(htmlcode, height=500, scrolling=True)
        else:
            st.warning("마인드맵이 생성되지 않았습니다.")
    
    st.write("")
    st.subheader("💾 내보내기 및 공유")
    
    col1, col2, col3, col4, col5 = st.columns(5)
    
    # [수정] 다운로드 파일에 들어갈 전체 내용을 재조립합니다 (구분선 제외)
    export_full_text = summary_text + "\n\n" + transcript_text + "\n\n" + mindmap_text
    
    doc = Document()
    doc.add_heading(f"[{event_type}] {event_title}", 0)
    for line in export_full_text.split('\n'):
        clean_line = line.strip()
        if not clean_line or clean_line == "---":
            continue
        if clean_line.startswith('# '):
            doc.add_heading(clean_line.replace('# ', ''), level=1)
        elif clean_line.startswith('## '):
            doc.add_heading(clean_line.replace('## ', ''), level=2)
        elif clean_line.startswith('### '):
            doc.add_heading(clean_line.replace('### ', ''), level=3)
        elif clean_line.startswith('- '):
            doc.add_paragraph(clean_line[2:], style='List Bullet')
        elif clean_line.startswith('**') and clean_line.endswith('**'):
            p = doc.add_paragraph()
            p.add_run(clean_line.replace('**', '')).bold = True
        else:
            doc.add_paragraph(clean_line)
            
    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)
    
    with col1:
        st.download_button(
            label="📥 Word",
            data=doc_stream,
            file_name=f"{event_date}_{event_title}_회의록.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
    with col2:
        st.download_button(
            label="📝 Markdown",
            data=export_full_text,
            file_name=f"{event_date}_{event_title}_회의록.md",
            mime="text/markdown"
        )
        
    with col3:
        html_content = markdown.markdown(export_full_text, extensions=['tables'])
        full_html = f"""<html>
<head>
<meta charset='utf-8'>
<style>
    body {{ font-family: 'Malgun Gothic', sans-serif; line-height: 1.6; padding: 20px; }}
    table {{ border-collapse: collapse; width: 100%; margin-bottom: 20px; }}
    th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
    th {{ background-color: #f2f2f2; }}
</style>
</head>
<body>{html_content}</body>
</html>"""
        st.download_button(
            label="🌐 HTML (PDF용)",
            data=full_html,
            file_name=f"{event_date}_{event_title}_회의록.html",
            mime="text/html"
        )
        
    with col4:
        # [수정] 대화 원본만 분리해서 다운로드 (sections[1] 활용)
        clean_transcript = transcript_text.replace("```", "")
            
        st.download_button(
            label="💬 대화록 TXT",
            data=clean_transcript,
            file_name=f"{event_date}_{event_title}_대화록.txt",
            mime="text/plain"
        )

    with col5:
        subject = urllib.parse.quote(f"[{event_type}] {event_title} 회의록 공유")
        body = urllib.parse.quote(export_full_text)
        mailto_link = f"mailto:?subject={subject}&body={body}"
        st.markdown(f'<a href="{mailto_link}"><button style="width:100%; border-radius:8px; padding:0.4rem; background-color:white; border:1px solid #dcdede; cursor:pointer;">📧 공유</button></a>', unsafe_allow_html=True)
        
    st.divider()
    with st.expander("🔄 2차 가공 (보고서/기획안 형식으로 자동 변환)", expanded=False):
        st.info("원하는 문서 양식이나 추가 지시사항을 입력하면 AI가 회의록을 바탕으로 새로운 문서를 만들어줍니다.\n(예: '주간 회의 보고서 형태로 만들어줘')")
        
        custom_prompt = st.text_area("추가 작업 지시 (Prompt):", placeholder="예: 내용을 요약해서 표로 만들어줘")
        
        col_btn1, col_btn2 = st.columns([1, 1])
        with col_btn1:
            if st.button("✨ 변환 실행"):
                if custom_prompt:
                    with st.spinner("AI가 문서를 재구성하고 있습니다..."):
                        try:
                            refine_prompt = f"""
다음은 회의록 원본 데이터입니다:

{st.session_state.first_pass_result}

위 내용을 바탕으로, 다음 사용자의 추가 지시사항에 맞게 문서를 완벽하게 재구성해 주세요:
[지시사항]: "{custom_prompt}"

전문적이고 깔끔한 비즈니스 문서 형태로 마크다운(표, 리스트 등)을 적극 활용하여 작성해 주세요.
반드시 앞서 사용한 `<!-- SECTION_DIVIDER -->` 구분선도 동일하게 유지하여 문서의 구조를 살려주세요.
"""
                            refined_response = model.generate_content([refine_prompt])
                            st.session_state.post_process_result = refined_response.text
                            st.rerun() 
                        except Exception as e:
                            st.error(f"변환 중 오류가 발생했습니다: {e}")
                else:
                    st.warning("지시사항을 입력해 주세요.")
        
        with col_btn2:
            if st.session_state.post_process_result:
                if st.button("↩️ 원본 회의록으로 복구"):
                    st.session_state.post_process_result = None
                    st.rerun()
